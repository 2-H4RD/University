"""Plot local cubic Lagrange interpolation for the first 56 observed weeks."""

from pathlib import Path
import matplotlib.pyplot as plt
import numpy as np
import openpyxl


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Практика 4.xlsx"
if not SOURCE.exists():
    SOURCE = HERE.parent / "Практика 3" / "Практика 3.xlsx"

book = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
sheet = book["Лист2"] if "Лист2" in book.sheetnames else book.active
y = [sheet.cell(row, 2).value for row in range(2, 58)]
if len(y) != 56 or not all(isinstance(v, (int, float)) for v in y):
    raise ValueError("Expected 56 numeric observations in column B, rows 2:57")

def lagrange_local(t):
    """Use four adjacent observed nodes around the containing interval."""
    if not 1 <= t <= 56:
        raise ValueError("Only the observation period 1..56 is plotted")
    if t == 56:
        return float(y[-1])
    interval = int(np.floor(t))
    first = min(max(interval - 1, 1), 53)
    support = range(first, first + 4)
    result = 0.0
    for i in support:
        basis = 1.0
        for j in support:
            if i != j:
                basis *= (t - j) / (i - j)
        result += y[i - 1] * basis
    return result


xgrid = np.linspace(1, 56, 1101)
pgrid = np.array([lagrange_local(t) for t in xgrid])
nodes = np.arange(1, 57)
assert all(abs(lagrange_local(i) - y[i - 1]) < 1e-12 for i in nodes)

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11})
fig, ax = plt.subplots(figsize=(12, 5.6), constrained_layout=True)
ax.plot(xgrid, pgrid, color="#bb493f", lw=1.65, label="Локальная интерполяция Лагранжа")
ax.scatter(nodes, y, color="#145ea8", s=24, zorder=3, label="Фактические значения")
ax.set_xlim(0.5, 56.5)
ax.set_ylim(min(0, pgrid.min() - 1), max(y) + 2)
ax.set_xticks([1, *range(5, 56, 5), 56])
ax.set_xlabel("Номер недели t")
ax.set_ylabel("Количество инцидентов")
ax.set_title("Интерполяция Лагранжа, недели 1–56")
ax.grid(True, color="#dddddd", linewidth=0.6)
ax.legend(loc="upper left", frameon=False)

output = HERE / "Полином Лагранжа 1-56.png"
fig.savefig(output, dpi=180)
print(f"source={SOURCE}")
print(f"output={output}")
print(f"samples_min={pgrid.min():.9g}, samples_max={pgrid.max():.9g}")
